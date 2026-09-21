"""Macros de pago masivo del BCP, con una plantilla y una base de cuentas
ficticias (misma estructura que las del banco, sin datos reales)."""
import io
import zipfile
from datetime import date

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from app.core.security import hash_password
from app.models.user import User
from app.services import macro_bcp
from app.services.excel_utils import ProcesamientoError

AMARILLO = "FFFFFF00"
FILA_CHEQUE = 20


def _plantilla(restos: bool = False) -> bytes:
    """Una .xlsm mínima: hoja de abonos con la fila 7, las filas modelo 11 (A,
    amarilla) y 12 (D), la sección de cheques en la 20, la macro y un botón."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = macro_bcp.HOJA_INPUT
    ws["A6"], ws["B6"] = "Tipo de Registro", "Cantidad de abonos de la planilla"
    ws["A7"], ws["D7"], ws["E7"], ws["G7"] = "C", "C", "1930000000026", "Referencia PaP"
    for c in range(1, 16):
        ws.cell(11, c).fill = PatternFill("solid", fgColor=AMARILLO)
        ws.cell(11, c).number_format = "@"
        ws.cell(12, c).number_format = "@"
    if restos:  # datos de una semana anterior
        ws["A15"], ws["G15"], ws["I15"] = "A", "PROVEEDOR VIEJO", "999.99"
    ws.cell(FILA_CHEQUE, 1).value = "DATOS DEL ABONO\nCON CHEQUE DE GERENCIA"
    ws.cell(FILA_CHEQUE + 1, 1).value = "Tipo de Registro"
    wb.create_sheet("Hoja1")["A1"] = "5173.97"
    buf = io.BytesIO()
    wb.save(buf)
    # openpyxl no escribe macros ni botones: se agregan al zip a mano.
    salida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin, \
            zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("xl/vbaProject.bin", b"macro del banco")
        zout.writestr("xl/ctrlProps/ctrlProp1.xml", b"<formControlPr/>")
    return salida.getvalue()


def _base(cabecera_ok: bool = True) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Tipo de Registro", "Tipo de Cuenta de Abono",
        "Cuenta de Abono" if cabecera_ok else "Otra cosa",
        "Tipo de Documento de Identidad", "Número de Documento de Identidad",
        "Correlativo", "Nombre del proveedor",
    ])
    ws.append(["A", "C", "1910000000011", "6", "20111111111", "   ", "PROVEEDOR UNO SAC"])
    # RUC guardado como número: debe normalizarse igual.
    ws.append(["A", "B", "00219300000000000022", 6, 20222222222.0, None, "PROVEEDOR DOS SA"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hoja(contenido: bytes):
    return openpyxl.load_workbook(io.BytesIO(contenido))[macro_bcp.HOJA_INPUT]


def _factura(ruc, numero, proveedor="DEL INFORME"):
    return {"RUC": ruc, "NUMERO": numero, "PROVEEDOR": proveedor}


@pytest.mark.parametrize("numero, esperado", [
    ("F002-00012369", "12369"),
    ("E001-6", "6"),
    ("FF2A-00173430", "173430"),
    ("F001-000", "0"),
    ("SIN-NUMERO", "NUMERO"),
])
def test_numero_documento(numero, esperado):
    assert macro_bcp.numero_documento(numero) == esperado


def test_leer_base_cuentas():
    cuentas = macro_bcp.leer_base_cuentas(_base())
    assert cuentas["20111111111"] == {
        "tipo_cuenta": "C", "cuenta": "1910000000011", "tipo_doc": "6",
        "nombre": "PROVEEDOR UNO SAC",
    }
    assert cuentas["20222222222"]["tipo_doc"] == "6"
    with pytest.raises(ProcesamientoError):
        macro_bcp.leer_base_cuentas(_base(cabecera_ok=False))


def test_armar_abonos_agrupa_por_ruc_y_cruza_la_base():
    cuentas = macro_bcp.leer_base_cuentas(_base())
    abonos = macro_bcp.armar_abonos(
        [
            (_factura("20111111111", "F001-00000010"), 100.004),
            (_factura("20999999999", "E001-7", "SIN CUENTA SAC"), 50),
            (_factura("20111111111", "F001-00000011"), 200.555),
        ],
        cuentas,
    )
    uno, sin = abonos
    # Nombre de la base, facturas redondeadas a 2 decimales y luego sumadas.
    assert (uno.nombre, uno.cuenta, uno.total) == ("PROVEEDOR UNO SAC", "1910000000011", 300.56)
    assert [d.numero for d in uno.documentos] == ["10", "11"]
    # Sin cuenta: va igual, con el nombre del informe y la cuenta en blanco.
    assert (sin.en_bd, sin.nombre, sin.cuenta) == (False, "SIN CUENTA SAC", "")


def test_generar_macro_llena_y_deja_lo_demas_intacto():
    plantilla = _plantilla()
    abonos = macro_bcp.armar_abonos(
        [
            (_factura("20111111111", "F001-00000010"), 1000),
            (_factura("20111111111", "F001-00000011"), 234.5),
            (_factura("20999999999", "E001-7", "SIN CUENTA SAC"), 50),
        ],
        macro_bcp.leer_base_cuentas(_base()),
    )
    salida = macro_bcp.generar_macro(plantilla, abonos, "USD", date(2026, 9, 15))
    ws = _hoja(salida)

    assert (ws["B7"].value, ws["C7"].value, ws["F7"].value) == ("000002", "20260915", "1284.50")
    assert ws["E7"].value == "1930000000026"          # cuenta de cargo: la de la plantilla
    fila = lambda r: [ws.cell(r, c).value for c in range(1, 16)]
    assert fila(11) == [
        "A", "C", "1910000000011", "6", "20111111111 ", "   ", "PROVEEDOR UNO SAC",
        "D", "1234.50", "S", "0002", None, None, None, None,
    ]
    assert fila(12)[0] == "D" and fila(12)[11:] == ["F ", "10", "D", "1000.00"]
    assert fila(13)[11:] == ["F ", "11", "D", "234.50"]
    assert fila(14)[:9] == ["A", None, None, "6", "20999999999 ", "   ", "SIN CUENTA SAC", "D", "50.00"]
    # Las filas de abono llevan el formato de la fila 11 de la plantilla.
    assert ws.cell(11, 1).fill.fgColor.rgb == ws.cell(14, 1).fill.fgColor.rgb == AMARILLO
    assert ws.cell(12, 1).fill.fgColor.rgb != AMARILLO
    # La sección de cheques sigue en su sitio.
    assert "CHEQUE" in ws.cell(FILA_CHEQUE, 1).value

    zt, zs = zipfile.ZipFile(io.BytesIO(plantilla)), zipfile.ZipFile(io.BytesIO(salida))
    cambiadas = [n for n in zt.namelist() if zt.read(n) != zs.read(n)]
    assert cambiadas == ["xl/worksheets/sheet1.xml"]   # macro y botón, intactos


def test_generar_macro_limpia_restos_de_otra_semana():
    abonos = macro_bcp.armar_abonos(
        [(_factura("20111111111", "F001-1"), 10)], macro_bcp.leer_base_cuentas(_base())
    )
    ws = _hoja(macro_bcp.generar_macro(_plantilla(restos=True), abonos, "SOL", date(2026, 9, 15)))
    assert ws["A15"].value is None and ws["G15"].value is None


def test_generar_macro_sin_espacio_suficiente():
    abonos = macro_bcp.armar_abonos(
        [(_factura(f"20{i:09d}", f"F001-{i}"), 1) for i in range(10)], {}
    )
    with pytest.raises(ProcesamientoError, match="solo admite"):
        macro_bcp.generar_macro(_plantilla(), abonos, "SOL", date(2026, 9, 15))


def test_plantilla_sin_macro_rechazada():
    with pytest.raises(ProcesamientoError, match="xlsm"):
        macro_bcp.validar_plantilla(_base())


def _auth_headers(client) -> dict:
    from app.api.deps import get_db

    gen = client.app.dependency_overrides[get_db]()
    db = next(gen)
    db.add(User(username="tester", hashed_password=hash_password("s3cret"), is_admin=True))
    db.commit()
    try:
        next(gen)
    except StopIteration:
        pass
    resp = client.post("/api/v1/auth/login", json={"username": "tester", "password": "s3cret"})
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


def test_subir_archivos(client):
    headers = _auth_headers(client)
    url = "/api/v1/macros/archivos"
    assert [a["nombre"] for a in client.get(url, headers=headers).json()] == [None] * 4

    r = client.post(f"{url}/plantilla_SOL", headers=headers, files={"archivo": ("m.xlsm", _plantilla())})
    assert r.status_code == 200 and r.json()["detalle"] == f"admite {FILA_CHEQUE - 11} filas"
    r = client.post(f"{url}/bd_SOL", headers=headers, files={"archivo": ("bd.xlsx", _base())})
    assert r.status_code == 200 and r.json()["detalle"] == "2 cuentas"
    # Una base donde va la plantilla se rechaza con un mensaje claro.
    r = client.post(f"{url}/plantilla_USD", headers=headers, files={"archivo": ("bd.xlsx", _base())})
    assert r.status_code == 422
    r = client.post(f"{url}/otro", headers=headers, files={"archivo": ("x", b"x")})
    assert r.status_code == 422


def test_seleccionar_facturas_suma_agentes_identificados():
    """Pago masivo por proveedor y, después, los pagos a agentes como en el
    Detalle: una fila por O/C con su Neto y la O/C como número de documento. Las
    O/C de un agente van en su abono; las sin agente, cada una en el suyo, con
    el RUC y la cuenta en blanco."""
    from app.services import detalle_export
    from app.services.macro_service import seleccionar_facturas

    def fila(ruc, prov, numero, pos, moneda="USD", oc="", saldo="100"):
        return {"RUC": ruc, "PROVEEDOR": prov, "NUMERO": numero, "__pos": pos,
                "MONEDA": moneda, "ORD_COMPRA": oc, "TIPO": "01", "IMPORTE": saldo,
                "PAGADO": "0", "SALDO": saldo, "DETRACCION": "0"}

    data = {
        "operaciones": [
            {"pos": 1, "texto": "Pago masivo proveedores", "moneda": "SOL"},
            {"pos": 2, "texto": "Pago masivo proveedores", "moneda": "USD"},
            {"pos": 3, "texto": "Materia Prima Exterior", "moneda": "USD"},
        ],
        "filas": [
            fila("20111111111", "PROVEEDOR UNO", "F001-1", 2),
            fila("20333333333", "OTRA OPERACION", "F001-9", 3),
            # O/C con agente: su factura y la de la naviera van al agente.
            fila("20213635531", "DOGANA S.A.", "F003-153142", 2, oc="10031696", saldo="394.35"),
            fila("20492185087", "HAPAG LLOYD", "F001-77", 2, oc="10031696", saldo="50"),
            fila("20213635531", "DOGANA S.A.", "F003-153160", 2, oc="32053-5", saldo="854.58"),
            # O/C de un proveedor relacionado sin factura del agente.
            fila("831197135", "NOURYON LLC", "5103572506", 3, oc="31637-4", saldo="90280.99"),
            fila("831197135", "NOURYON LLC", "5103572505", 3, oc="31637-5", saldo="30104.41"),
            fila("831197135", "NOURYON LLC", "5103572507", 3, oc="30959-A", saldo="10"),
        ],
    }
    calc = detalle_export.preparar_calculo(
        data, agente_rucs=["20213635531"], relacionados_rucs=["831197135"]
    )
    facturas = seleccionar_facturas(data, calc)

    abonos = macro_bcp.armar_abonos(facturas["USD"], {})
    assert [(a.ruc, a.agente, a.total) for a in abonos] == [
        ("20111111111", False, 100.0),
        ("20213635531", True, 1298.93),  # sus dos O/C en un solo abono
        ("", True, 90280.99),            # sin agente: un abono por O/C
        ("", True, 30104.41),
        ("", True, 10.0),
    ]
    # Cada O/C es un documento, con el número de O/C sin guion y su Neto
    # (en la 10031696, 394.35 del agente + 50 de la naviera).
    assert [(d.numero, d.monto) for d in abonos[1].documentos] == [
        ("10031696", 444.35), ("320535", 854.58),
    ]
    assert [d.numero for d in abonos[2].documentos] == ["316374"]
    assert [d.numero for d in abonos[4].documentos] == ["30959"]   # sin la letra
    assert abonos[2].nombre == "Colocar nombre de agente manualmente"
    assert (abonos[2].cuenta, abonos[2].en_bd) == ("", False)
    assert facturas["SOL"] == []
